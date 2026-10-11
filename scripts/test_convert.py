"""Local check of api/convert.py without Blob: builds sample files, serves them via a stubbed download.
Run: .venv/bin/python scripts/test_convert.py"""
import os, shutil, sys, tempfile
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "api"))
import convert as c
from docx import Document
from reportlab.pdfgen import canvas

d = tempfile.mkdtemp()
pdf, docx, csv, png = (os.path.join(d, n) for n in ("a.pdf", "b.docx", "c.csv", "d.png"))
k = canvas.Canvas(pdf); k.drawString(72, 720, "Hello PDF world"); k.save()
doc = Document(); doc.add_heading("Title Here", 1); doc.add_paragraph("Body text."); doc.save(docx)
open(csv, "w").write("x,y\n1,2\n"); open(png, "wb").write(b"x")

deleted = []
c.check_url = lambda u: None
c.delete_blob = deleted.append
c.download = lambda url, path: shutil.copy(url, path)

for f, expect in ((pdf, "Hello PDF world"), (docx, "# Title Here"), (csv, "| x | y |")):
    out = c.convert(f, os.path.basename(f))
    assert expect in out, (f, out)
    print("ok", os.path.basename(f), repr(out[:40]))
assert len(deleted) == 3
try:
    c.convert(png, "d.png"); raise SystemExit("png should be rejected")
except c.ConvertError as e:
    print("ok rejected png:", e)
assert not [x for x in os.listdir("/tmp") if x.endswith((".pdf", ".docx", ".csv")) and x.startswith("tmp")], "temp leak"
# SSRF guard
import importlib; importlib.reload(c)
try:
    c.check_url("https://evil.example.com/x.pdf"); raise SystemExit("ssrf not blocked")
except c.ConvertError:
    print("ok ssrf blocked")
